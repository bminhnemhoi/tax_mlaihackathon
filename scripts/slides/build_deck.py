# -*- coding: utf-8 -*-
"""Dựng bộ slide trình bày vòng cải tiến 17/10 của Tax Referee (16:9, python-pptx).

Chạy lại bao nhiêu lần cũng được:
    python scripts/slides/build_deck.py
    python scripts/slides/build_deck.py --screens docs/submission/screens --out docs/submission/x.pptx

Nếu thư mục ảnh có dashboard.png, batch.png, explain.png, channels.png, history.png, trust.png
thì ảnh được chèn vào slide tương ứng; thiếu ảnh nào thì vẽ khung chờ có nhãn.
Màu và phông theo bộ 5 slide của đội (Cambria cho tiêu đề, Calibri cho chữ).
"""
from __future__ import annotations

import argparse
import os
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

try:
    from PIL import Image
except ImportError:  # Pillow chỉ cần để giữ tỉ lệ ảnh
    Image = None

ROOT = Path(__file__).resolve().parents[2]

# Bảng màu lấy từ theme của bộ 5 slide (Tax-Referee_5-slide_trinh-bay-17-10.pptx)
NAVY = "0F172A"
CARD_DARK = "1E293B"
SLATE = "334155"
MUTED = "64748B"
SOFT = "94A3B8"
LINE = "CBD5E1"
ON_DARK = "E2E8F0"
LIGHT = "F1F5F9"
PAPER = "F8FAFC"
WHITE = "FFFFFF"
LIME = "D9F99D"
LIME_DEEP = "4D7C0F"
LIME_TINT = "ECFCCB"
AMBER = "B45309"
AMBER_TINT = "FEF3C7"
ROSE = "BE123C"
ROSE_TINT = "FFE4E6"
TEAL = "0F766E"

HEAD = "Cambria"
BODY = "Calibri"
FOOTER = "Tax Referee · Đội Fprinciples · MLAI Hackathon 2026 · Track VNG"
LIVE = "limegreen-salamander-517622.hostingersite.com"


NBSP_UNITS = re.compile(r"(\d) (triệu|giây|ms|ngày|người|ca|bộ|lần|trường|phép|thao)")


def keep_units(t: str) -> str:
    """Không để số và đơn vị rơi xuống hai dòng khác nhau ("5" / "triệu")."""
    return NBSP_UNITS.sub(lambda m: m.group(1) + chr(0xA0) + m.group(2), t)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- văn bản
def R(text, **o):
    return ("run", text, o)


def P(*runs, **o):
    return ("para", runs, o)


def _bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.2))))
    pPr.set("indent", str(-int(Inches(0.2))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def fill_tf(tf, paras, **base):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, para in enumerate(paras):
        if isinstance(para, str):
            para = P(para)
        _, runs, po = para
        opts = {**base, **po}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        if opts.get("before"):
            p.space_before = Pt(opts["before"])
        if opts.get("after") is not None:
            p.space_after = Pt(opts["after"])
        if opts.get("line"):
            p.line_spacing = opts["line"]
        if opts.get("bullet"):
            _bullet(p, opts.get("bullet_color", LIME_DEEP))
        for r in runs:
            if isinstance(r, str):
                r = R(r)
            _, text, ro = r
            ro = {**opts, **ro}
            run = p.add_run()
            run.text = keep_units(text)
            f = run.font
            f.size = Pt(ro.get("size", 14))
            f.bold = bool(ro.get("bold", False))
            f.italic = bool(ro.get("italic", False))
            f.name = ro.get("font", BODY)
            f.color.rgb = rgb(ro.get("color", SLATE))


def _no_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, **base):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    fill_tf(tf, paras, **base)
    return s


def box(slide, x, y, w, h, fill=LIGHT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        paras=None, anchor=MSO_ANCHOR.MIDDLE, pad=0.12, dash=False, line_w=1.0, name=None, **base):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _no_style(s)
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(min(pad, 0.06))
    tf.vertical_anchor = anchor
    if paras is not None:
        if isinstance(paras, (str, tuple)):
            paras = [paras]
        fill_tf(tf, paras, **base)
    return s


def circle(slide, x, y, d, label, fill=NAVY, color=WHITE, size=12):
    return box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, pad=0,
               paras=[P(R(label, bold=True), align=PP_ALIGN.CENTER)], size=size, color=color)


def arrow(slide, x1, y1, x2, y2, color=SOFT, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _no_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


# ---------------------------------------------------------------- khung slide
class Deck:
    def __init__(self, screens: Path):
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(5.625)
        self.screens = screens
        self.n = 0
        self.placeholders: list[str] = []
        self.inserted: list[str] = []

    def slide(self, dark=False, kicker=None, title=None, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
        if kicker:
            text(s, 0.5, 0.24, 9.0, 0.25, kicker, size=10, bold=True,
                 color=LIME if dark else LIME_DEEP, name="Kicker")
        if title:
            text(s, 0.5, 0.48, 9.0, 0.6, title, size=26, bold=True, font=HEAD,
                 color=WHITE if dark else NAVY, anchor=MSO_ANCHOR.TOP, name="Title")
        if self.n > 1:
            text(s, 0.5, 5.2, 7.6, 0.25, FOOTER, size=9, color=SOFT if dark else SLATE, name="Footer")
            text(s, 8.9, 5.2, 0.6, 0.25, P(str(self.n), align=PP_ALIGN.RIGHT), size=9,
                 color=SOFT if dark else SLATE, name="Slide number")
        if notes:
            s.notes_slide.notes_text_frame.text = " ".join(notes.split())
        return s

    def screen(self, s, key, label, x, y, w, h, align="right"):
        path = self.screens / f"{key}.png"
        if path.is_file():
            if Image is not None:
                with Image.open(path) as im:
                    iw, ih = im.size
                sc = min(w / iw, h / ih)
                pw, ph = iw * sc, ih * sc
            else:
                pw, ph = w, h
            # bám mép ngoài của cột và mép trên để các slide thẳng hàng với nhau
            px = x + (w - pw) if align == "right" else x
            py = y
            pic = s.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
            pic.name = f"Ảnh màn hình {key}"
            pic.line.color.rgb = rgb(LINE)
            pic.line.width = Pt(0.75)
            self.inserted.append(key)
        else:
            box(s, x, y, w, h, fill=PAPER, line=SOFT, dash=True, radius=0.04, name=f"Khung chờ {key}",
                paras=[P(R(f"Ảnh màn hình: {label}", bold=True), align=PP_ALIGN.CENTER, after=4),
                       P(R(f"docs/submission/screens/{key}.png", size=9, color=SOFT), align=PP_ALIGN.CENTER)],
                size=13, color=MUTED)
            self.placeholders.append(key)


# ---------------------------------------------------------------- nội dung
def build(screens: Path, out: Path):
    d = Deck(screens)

    # 1 ─ Bìa
    s = d.slide(dark=True, notes="""
Chào ban giám khảo. Chúng em là đội Fprinciples, sản phẩm là Tax Referee, một trọng tài cho thuế giá trị gia tăng
đầu vào trên hóa đơn điện tử. Cả bài hôm nay gói trong một câu: máy quyết từng hóa đơn, còn người chốt cả lô.
Quy tắc ra quyết định, AI chỉ làm việc đọc chứng từ. Em sẽ đi qua cách làm theo lô mới thêm trong vòng này,
trả lời hai câu hỏi của Ban tổ chức về bảo mật và kiến trúc, rồi demo trực tiếp trên bản đang chạy.""")
    text(s, 0.6, 0.55, 8.8, 0.3, "TAX REFEREE · ĐỀ A — THE ESCALATION REFEREE · VÒNG CẢI TIẾN 17/10",
         size=11, bold=True, color=LIME)
    text(s, 0.6, 0.95, 8.8, 1.45, [P("Máy quyết từng hóa đơn."), P(R("Người chốt cả lô.", color=LIME))],
         size=40, bold=True, font=HEAD, color=WHITE)
    text(s, 0.6, 2.5, 8.8, 0.4, "Trọng tài thuế GTGT đầu vào cho hóa đơn điện tử Việt Nam", size=18, color=ON_DARK)
    text(s, 0.6, 2.92, 8.8, 0.35, "Quy tắc tất định quyết định. AI chỉ đọc chứng từ.", size=15, color=LINE)
    for i, (num, lab) in enumerate([("0", "lần gọi LLM trên đường quyết định"),
                                    ("9/9", "ca Verify đạt (4 ca chung + 5 ca Đề A)"),
                                    ("12/12", "phép thử bảo mật trên bản đang chạy")]):
        x = 0.6 + i * 2.95
        box(s, x, 3.55, 2.75, 1.05, fill=CARD_DARK, radius=0.1)
        text(s, x + 0.2, 3.62, 2.4, 0.5, num, size=26, bold=True, font=HEAD, color=LIME)
        text(s, x + 0.2, 4.12, 2.4, 0.42, lab, size=11, color=LINE)
    text(s, 0.6, 5.0, 8.8, 0.3, f"Đội Fprinciples · MLAI Hackathon 2026 · Track VNG · {LIVE}", size=10, color=SOFT)

    # 2 ─ Vấn đề
    s = d.slide(kicker="VẤN ĐỀ", title="Cuối tháng: hàng trăm hóa đơn, một chị kế toán", notes="""
Thử hình dung chị kế toán vào cuối tháng. Hóa đơn đổ về từ mail, từ Zalo, có cái là file XML, có cái chỉ là ảnh chụp.
Mỗi hóa đơn phải đạt nhiều điều kiện cùng lúc, và chỉ cần sai một điều kiện là mất tiền thuế được khấu trừ.
Ví dụ một hóa đơn sáu triệu trả bằng tiền mặt, lập sau ngày 1 tháng 7 năm 2025, thì không được khấu trừ,
vì ngưỡng bây giờ là năm triệu. Có hai cách làm sai. Một là cho qua hết: nhanh, nhưng khi thanh tra thì bị loại.
Hai là cái gì cũng hỏi sếp: an toàn, nhưng sếp nghẽn, và tự động hóa không còn ý nghĩa.""")
    for i, ch in enumerate(["Gmail", "Zalo", "Tải lên"]):
        y = 1.4 + i * 0.55
        box(s, 0.5, y, 1.1, 0.4, fill=WHITE, line=LINE, paras=[P(ch, align=PP_ALIGN.CENTER)], size=12,
            bold=True, color=SLATE, pad=0.04)
        arrow(s, 1.65, y + 0.2, 2.05, 2.0)
    box(s, 2.1, 1.4, 2.5, 1.2, fill=LIGHT, anchor=MSO_ANCHOR.MIDDLE,
        paras=[P(R("Hàng trăm hóa đơn", bold=True, size=15, color=NAVY), align=PP_ALIGN.CENTER),
               P("XML, PDF, ảnh chụp", align=PP_ALIGN.CENTER, size=12, color=MUTED)])
    text(s, 0.5, 3.0, 4.1, 0.8, "Mỗi hóa đơn phải đạt nhiều điều kiện cùng lúc. Sai một điều kiện là mất tiền thuế được khấu trừ.",
         size=14, color=SLATE)
    box(s, 0.5, 3.9, 4.1, 0.95, fill=LIME_TINT, anchor=MSO_ANCHOR.MIDDLE, pad=0.15,
        paras=[P(R("Ví dụ: ", bold=True), "hóa đơn 6 triệu trả tiền mặt, lập sau 01/7/2025, không được khấu trừ. Ngưỡng nay là 5 triệu.")],
        size=12.5, color=NAVY)
    text(s, 5.0, 1.4, 4.5, 0.3, "Hai cách làm sai", size=14, bold=True, color=NAVY)
    for i, (t, b, fill, c) in enumerate([("Cho qua hết", "Nhanh. Nhưng thanh tra loại thuế, còn bị phạt.", ROSE_TINT, ROSE),
                                         ("Cái gì cũng hỏi sếp", "An toàn. Nhưng sếp nghẽn cuối tháng, tự động hóa thành vô nghĩa.", AMBER_TINT, AMBER)]):
        y = 1.85 + i * 1.55
        box(s, 5.0, y, 4.5, 1.35, fill=fill, radius=0.08)
        circle(s, 5.2, y + 0.42, 0.5, str(i + 1), fill=c, size=14)
        text(s, 5.9, y + 0.22, 3.45, 0.35, t, size=16, bold=True, color=NAVY)
        text(s, 5.9, y + 0.6, 3.45, 0.65, b, size=13, color=SLATE)

    # 3 ─ Ý tưởng
    s = d.slide(kicker="Ý TƯỞNG", title="Quy tắc quyết định, AI chỉ đọc chứng từ", notes="""
Ý tưởng của chúng em giống cô giáo chấm bài. Bài đúng hết thì cho điểm luôn. Bài sai chỗ nào thì khoanh đỏ
đúng chỗ đó và hỏi lại. Tax Referee làm y như vậy: hóa đơn đạt hết điều kiện thì máy tự duyệt; chưa đạt thì dừng,
nói rõ dừng vì sao, và hỏi đúng người. Chỉ có ba kiểu dừng: chưa rõ thông tin, ngoài quy định, và vượt thẩm quyền.
Phần quyết định không gọi mô hình ngôn ngữ lần nào, nên cùng một hóa đơn chạy lại bao nhiêu lần cũng ra cùng kết quả.
Ca máy tự duyệt vẫn bị hậu kiểm ngẫu nhiên mười phần trăm.""")
    box(s, 0.5, 1.35, 4.3, 2.2, fill=LIGHT, anchor=MSO_ANCHOR.MIDDLE, pad=0.25,
        paras=[P(R("Giống cô giáo chấm bài", bold=True, size=13, color=LIME_DEEP), after=8),
               P(R("Bài đúng hết thì cho điểm luôn. Bài sai chỗ nào thì khoanh đỏ đúng chỗ đó và hỏi lại.",
                   font=HEAD, italic=True, size=17, color=NAVY))])
    text(s, 0.5, 3.8, 4.3, 1.0, [P(R("0 lần gọi LLM", bold=True, color=NAVY), " trên đường quyết định.", after=4),
                                 P("Cùng một hóa đơn, chạy lại bao nhiêu lần cũng ra một kết quả.")], size=13, color=SLATE)
    box(s, 5.2, 1.35, 4.3, 0.55, fill=NAVY, paras=[P("Hóa đơn đạt hết điều kiện?", align=PP_ALIGN.CENTER)],
        size=14, bold=True, color=WHITE)
    arrow(s, 6.2, 1.92, 6.2, 2.3)
    arrow(s, 8.5, 1.92, 8.5, 2.3)
    box(s, 5.2, 2.35, 2.0, 0.9, fill=LIME_TINT, paras=[P(R("Đạt", bold=True, color=LIME_DEEP), align=PP_ALIGN.CENTER),
                                                       P("Máy tự duyệt", align=PP_ALIGN.CENTER, bold=True)], size=13, color=NAVY)
    box(s, 7.5, 2.35, 2.0, 0.9, fill=ROSE_TINT, paras=[P(R("Chưa đạt", bold=True, color=ROSE), align=PP_ALIGN.CENTER),
                                                       P("Dừng, hỏi đúng người", align=PP_ALIGN.CENTER, bold=True)], size=13, color=NAVY)
    text(s, 5.2, 3.45, 2.0, 1.3, "Vẫn hậu kiểm ngẫu nhiên 10% ca máy duyệt, gấp 3 khi sát ngưỡng.", size=12, color=MUTED)
    for i, stop in enumerate(["Chưa rõ thông tin", "Ngoài quy định", "Vượt thẩm quyền"]):
        box(s, 7.5, 3.42 + i * 0.47, 2.0, 0.38, fill=WHITE, line=ROSE, line_w=0.75,
            paras=[P(stop, align=PP_ALIGN.CENTER)], size=12, color=ROSE, bold=True, pad=0.04)

    # 4 ─ Luồng mới
    s = d.slide(kicker="CÁCH LÀM MỚI", title="Từ lúc nhận hóa đơn đến lúc ghi sổ", notes="""
Đây là luồng làm việc mới. Bước một, thu hóa đơn từ Gmail, Zalo hoặc tải lên, và biết ngay ai gửi. Bước hai,
kiểm theo lô: xếp theo mã số thuế, ngày, số hóa đơn, rồi máy quyết từng hóa đơn. Bước ba, ca nào dừng thì hỏi đúng người,
trả lời một lần kiểu "1A, 2A". Bước bốn là điểm mới quan trọng nhất: Kế toán trưởng hoặc CFO phải chốt lô thì mới ghi sổ.
Bước năm, sếp xem tổng quan. Và mọi bước đều vào nhật ký chỉ ghi thêm. Giống dây chuyền kiểm hàng: máy soi từng món,
nhưng tổ trưởng ký phiếu thì xe hàng mới được xuất kho.""")
    steps = [("Thu", "Gmail, Zalo, tải lên. Biết ai gửi."),
             ("Kiểm theo lô", "Xếp MST → ngày → số. Máy quyết từng hóa đơn."),
             ("Hỏi đúng người", "Trả lời một lần: “1A, 2A”."),
             ("Người chốt lô", "KTT hoặc CFO ký thì mới ghi sổ."),
             ("Sếp theo dõi", "Tự xử lý bao nhiêu, chờ bao nhiêu, VAT."),
             ("Nhật ký", "Chỉ ghi thêm. Sửa là lộ.")]
    bw, gap = 1.32, 0.216
    for i, (t, desc) in enumerate(steps):
        x = 0.5 + i * (bw + gap)
        key = i == 3
        box(s, x, 1.45, bw, 1.0, fill=NAVY if key else LIGHT, anchor=MSO_ANCHOR.MIDDLE, pad=0.06,
            paras=[P(R(str(i + 1), size=11, color=LIME if key else LIME_DEEP, bold=True), align=PP_ALIGN.CENTER),
                   P(R(t, bold=True), align=PP_ALIGN.CENTER)], size=13, color=WHITE if key else NAVY)
        if i < 5:
            arrow(s, x + bw + 0.02, 1.95, x + bw + gap - 0.02, 1.95, width=1.25)
        text(s, x + 0.02, 2.6, bw - 0.04, 1.2, desc, size=12.5, color=SLATE)
    box(s, 0.5, 3.95, 9.0, 0.8, fill=LIME_TINT, anchor=MSO_ANCHOR.MIDDLE, pad=0.25,
        paras=[P(R("Giống dây chuyền kiểm hàng: ", bold=True),
                 "máy soi từng món, nhưng tổ trưởng ký phiếu thì xe hàng mới được xuất kho.")], size=14, color=NAVY)

    # 5 ─ Theo lô
    s = d.slide(kicker="MỚI TRONG VÒNG NÀY · THEO LÔ", title="Nhiều file một lần, thứ tự luôn như nhau", notes="""
Giờ chị kế toán tải cả lô một lần, hoặc hóa đơn tự vào từ Gmail và Zalo.
Hệ thống đọc hết rồi mới xếp: theo mã số thuế, rồi ngày, rồi số hóa đơn, sau đó mới kiểm. Giống xếp bài thi theo số báo danh
trước khi chấm, ai chấm lại cũng ra cùng một thứ tự. Một file gửi hai lần, trùng mã băm SHA-256, chỉ tính một lần.
Với lô lớn, những hóa đơn cùng một câu hỏi được trả lời chung một lần, nhưng từng hóa đơn vẫn được kiểm và ghi nhật ký riêng.""")
    rows = ["Tải nhiều file một lần, hoặc nhận từ Gmail, Zalo.",
            "Đọc hết rồi mới xếp: MST → ngày → số, sau đó mới kiểm.",
            "Một file gửi hai lần (trùng SHA-256) chỉ tính một."]
    for i, r in enumerate(rows):
        y = 1.4 + i * 0.72
        circle(s, 0.5, y + 0.04, 0.42, str(i + 1), fill=LIME_DEEP, size=13)
        text(s, 1.1, y, 3.6, 0.6, r, size=14, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.5, 3.7, 4.2, 1.2, fill=LIGHT, anchor=MSO_ANCHOR.MIDDLE, pad=0.2,
        paras=[P(R("Lô lớn: ", bold=True, color=NAVY),
                 "hóa đơn cùng một câu hỏi được trả lời chung một lần. Từng hóa đơn vẫn được kiểm và ghi nhật ký riêng.")],
        size=13, color=SLATE)
    d.screen(s, "batch", "Lô hóa đơn", 5.0, 1.35, 4.5, 3.55)

    # 6 ─ Phiếu giải trình
    s = d.slide(kicker="MỚI TRONG VÒNG NÀY · PHIẾU GIẢI TRÌNH", title="Vì sao máy duyệt, vì sao máy dừng", notes="""
Muốn sếp tin máy thì máy phải nói được lý do. Mỗi hóa đơn có một phiếu giải trình. Khi máy duyệt, phiếu liệt kê từng điều kiện
đã đạt, có con số đi kèm, ai đọc cũng tự kiểm lại được. Khi máy dừng, phiếu tô đỏ đúng quy tắc bị vướng, đúng trường dữ liệu,
và ghi rõ bao nhiêu tiền thuế đang bị treo. Giống bài kiểm tra trả về có lời phê: đúng chỗ nào, sai chỗ nào, mất bao nhiêu điểm.""")
    box(s, 0.5, 1.35, 4.2, 1.6, fill=LIME_TINT, anchor=MSO_ANCHOR.MIDDLE, pad=0.25,
        paras=[P(R("Khi máy duyệt", bold=True, size=16, color=LIME_DEEP), after=6),
               P("Liệt kê từng điều kiện đã đạt, có con số đi kèm. Ai đọc cũng tự kiểm lại được.")],
        size=13.5, color=NAVY)
    box(s, 0.5, 3.2, 4.2, 1.7, fill=ROSE_TINT, anchor=MSO_ANCHOR.MIDDLE, pad=0.25,
        paras=[P(R("Khi máy dừng", bold=True, size=16, color=ROSE), after=6),
               P("Tô đỏ đúng quy tắc, đúng trường dữ liệu, và số VAT đang bị treo.")],
        size=13.5, color=NAVY)
    d.screen(s, "explain", "Phiếu giải trình", 5.0, 1.35, 4.5, 3.55)

    # 7 ─ Máy quyết, người chốt
    s = d.slide(kicker="MỚI TRONG VÒNG NÀY · MÁY QUYẾT, NGƯỜI CHỐT", title="Ai làm gì trong một lô", notes="""
Nguyên tắc của vòng này là: máy quyết, người chốt. Máy quyết từng hóa đơn theo quy tắc và ghi nhật ký, nhưng không ghi sổ
khi lô chưa được chốt. Kế toán viên tạo lô và gắn hóa đơn. Kế toán trưởng duyệt ca dưới hai trăm triệu và chốt lô.
Ca từ hai trăm triệu trở lên là việc của CFO, và lô có hóa đơn như vậy thì chỉ CFO chốt được. Người tạo lô không tự chốt lô của mình,
để luôn có bốn mắt. Tóm lại: chưa có chữ ký chốt lô thì chưa có gì vào sổ.""")
    cols = [(0.5, 1.9, "Vai"), (2.4, 3.75, "Được làm"), (6.15, 3.35, "Không được làm")]
    box(s, 0.5, 1.3, 9.0, 0.42, fill=NAVY, radius=0.1)
    for x, w, h in cols:
        text(s, x + 0.15, 1.3, w - 0.2, 0.42, h, size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    table = [("Máy (trọng tài)", "Quyết từng hóa đơn theo quy tắc, ghi nhật ký", "Ghi sổ khi lô chưa được chốt"),
             ("Kế toán viên", "Tạo lô, gắn hóa đơn, xem kết quả", "Chốt lô"),
             ("Kế toán trưởng", "Duyệt ca dưới 200.000.000₫; chốt lô", "Chốt lô mình tạo; chốt lô có ca từ 200 triệu"),
             ("CFO", "Duyệt ca từ 200.000.000₫, chốt lô lớn", "Chốt lô mình tạo")]
    for i, row in enumerate(table):
        y = 1.78 + i * 0.6
        box(s, 0.5, y, 9.0, 0.55, fill=LIGHT if i % 2 == 0 else WHITE, radius=0.1)
        for (x, w, _), cell in zip(cols, row):
            text(s, x + 0.15, y, w - 0.25, 0.55, cell, size=12.5, color=NAVY if x == 0.5 else SLATE,
                 bold=(x == 0.5), anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.5, 4.3, 9.0, 0.6, fill=LIME_TINT, pad=0.25,
        paras=[P("Chưa có chữ ký chốt lô của Kế toán trưởng hoặc CFO thì chưa có gì vào sổ.", align=PP_ALIGN.CENTER)],
        size=15, bold=True, color=NAVY)

    # 8 ─ Người gửi
    s = d.slide(kicker="MỚI TRONG VÒNG NÀY · NGƯỜI GỬI", title="Biết ai gửi, theo dõi hóa đơn của từng người", notes="""
Hóa đơn đến qua mail thì phải biết ai gửi thật. Email chỉ được tính là của nhân viên khi địa chỉ có trong danh bạ và qua
kiểm tra xác thực SPF hoặc DKIM. Cùng địa chỉ đó mà không qua xác thực thì bị đánh dấu nghi giả mạo. Người ngoài danh bạ thì chờ kế toán
xác nhận. Giống bảo vệ tòa nhà: có thẻ đúng thì cho vào; mặc đồng phục mà không có thẻ thì giữ lại hỏi. Máy báo cho Kế toán trưởng,
trả lời người gửi qua Zalo hoặc Gmail, và ai cũng có trang Lịch sử của tôi để xem hóa đơn mình đã gửi, đã duyệt.""")
    tags = [("Nhân viên", LIME_TINT, LIME_DEEP, "Email trong danh bạ và qua SPF/DKIM, hoặc Zalo khớp danh bạ."),
            ("Nghi giả mạo", ROSE_TINT, ROSE, "Trùng địa chỉ nhân viên nhưng không qua xác thực."),
            ("Chờ xác nhận", AMBER_TINT, AMBER, "Ngoài danh bạ. Kế toán xác nhận rồi mới xử lý.")]
    for i, (t, fill, c, desc) in enumerate(tags):
        y = 1.4 + i * 0.78
        box(s, 0.5, y + 0.1, 1.3, 0.38, fill=fill, paras=[P(t, align=PP_ALIGN.CENTER)], size=12, bold=True, color=c, pad=0.03)
        text(s, 1.95, y, 2.85, 0.6, desc, size=12.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.5, 3.85, 4.3, 1.0, [P("Máy báo Kế toán trưởng, trả lời người gửi qua Zalo hoặc Gmail.", after=4),
                                  P(R("Ai cũng có trang “Lịch sử của tôi”.", bold=True, color=NAVY))], size=13, color=SLATE)
    d.screen(s, "channels", "Kênh Gmail/Zalo và người gửi", 5.1, 1.35, 4.4, 1.7)
    d.screen(s, "history", "Lịch sử của tôi", 5.1, 3.2, 4.4, 1.7)

    # 9 ─ Dashboard
    s = d.slide(kicker="MỚI TRONG VÒNG NÀY · CHO SẾP", title="Tổng quan cho Kế toán trưởng và CFO", notes="""
Sếp không cần mở từng hóa đơn. Màn hình tổng quan cho biết bao nhiêu phần trăm máy tự xử lý, bao nhiêu người duyệt,
bao nhiêu bị từ chối. Tiền thuế đã khấu trừ và tiền thuế đang chờ. Lý do dừng hay gặp nhất, để biết nên sửa quy trình ở đâu.
Câu hỏi được trả lời sau bao lâu. Và chia theo kênh, theo người gửi, để thấy ai hay gửi hóa đơn thiếu.""")
    metrics = ["% máy tự xử lý, người duyệt, từ chối", "VAT đã khấu trừ và VAT đang chờ", "Lý do dừng hay gặp nhất",
               "Câu hỏi được trả lời sau bao lâu", "Theo kênh, theo người gửi"]
    text(s, 0.5, 1.45, 3.4, 3.4, [P(m, bullet=True, after=12) for m in metrics], size=14, color=SLATE)
    d.screen(s, "dashboard", "Tổng quan cho sếp", 4.2, 1.3, 5.3, 3.65)

    # 10 ─ Q1: 4 trụ
    s = d.slide(kicker="CÂU HỎI 1 · BẢO MẬT HỆ CÓ AGENT AI TRONG DOANH NGHIỆP",
                title="Máy đọc có thể bị lừa, trọng tài thì không", notes="""
Câu hỏi một là bảo mật. Câu trả lời ngắn: chúng em cố ý không để AI quyết định. Định danh: máy tự duyệt thì nhật ký ghi tên
trọng tài kèm mã băm của bộ quy tắc và dữ liệu, máy không mượn danh người nhập. Đặc quyền tối thiểu: quyền kiểm ở server,
Kế toán trưởng sửa ca ba trăm ba mươi triệu xuống một trăm chín mươi tám triệu thì ca vẫn phải chờ CFO. Ranh giới tất định:
agent đọc ảnh chỉ được khai mười bốn trường. Cách ly: mỗi phiên một database, nhật ký bị sửa là lộ. Giới hạn em nói thẳng:
vai trò trong demo là mô phỏng, trình đọc PDF chưa nằm trong container riêng.""")
    pillars = [
        ("Định danh", "Máy tự duyệt ghi tên SYSTEM_REFEREE kèm hash policy, hash dữ liệu. Email chỉ là nhân viên khi qua SPF/DKIM.",
         "phép thử I1, I2", "vai trò trong demo là mô phỏng; production cần SSO, MFA cho CFO."),
        ("Đặc quyền tối thiểu", "Quyền kiểm ở server. KTT hạ ca 330 triệu xuống 198 triệu vẫn phải chờ CFO. Người tạo lô không tự chốt.",
         "L1, L2, L3", "“đã có ủy nhiệm chi” vẫn do kế toán khai; bù bằng hậu kiểm."),
        ("Ranh giới thực thi tất định", "0 lần gọi LLM khi quyết định. Agent đọc ảnh chỉ khai được 14 trường; workflow chỉ nhận 7 thao tác.",
         "D1 đến D4, /api/health", "agent bị lừa vẫn đọc sai được giá trị 14 trường, kể cả số tiền."),
        ("Cách ly", "Mỗi phiên một database riêng. Nhật ký chỉ ghi thêm, chuỗi SHA-256, HMAC, điểm neo.",
         "S1, S2, S3", "cùng tiến trình, khóa HMAC cùng máy; trình đọc PDF chưa ở container riêng."),
    ]
    for i, (name, mech, ev, lim) in enumerate(pillars):
        x = 0.5 + (i % 2) * 4.6
        y = 1.3 + (i // 2) * 1.85
        box(s, x, y, 4.4, 1.72, fill=LIGHT, radius=0.06)
        circle(s, x + 0.15, y + 0.13, 0.34, str(i + 1), fill=NAVY, size=11)
        text(s, x + 0.6, y + 0.1, 3.7, 0.4, name, size=14, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.2, y + 0.55, 4.05, 1.12,
             [P(R("Cơ chế: ", bold=True, color=NAVY), mech, after=2),
              P(R("Bằng chứng: ", bold=True, color=LIME_DEEP), ev, after=2),
              P(R("Giới hạn: ", bold=True, color=AMBER), lim)], size=10.5, color=SLATE)

    # 11 ─ /trust
    s = d.slide(kicker="CÂU HỎI 1 · BẰNG CHỨNG", title="Bằng chứng chạy trên chính bản đang deploy", notes="""
Nói bảo mật thì phải chứng minh được. Trang trust có nút chạy mười hai phép thử tấn công, chia theo bốn trụ vừa nói,
chạy ngay trên mã đang deploy chứ không phải ảnh chụp. Trang health cho thấy số lần gọi mô hình ngôn ngữ trên đường quyết định
bằng không. Ngoài ra bộ bảy mươi hai ca red-team đều đạt. Chỗ nào chưa chặn được, như agent đọc sai số tiền,
chúng em ghi thẳng trong phần giới hạn.""")
    d.screen(s, "trust", "Trang /trust: 12 phép thử", 0.5, 1.3, 5.4, 3.65, align="left")
    for i, (num, lab) in enumerate([("12/12", "phép thử tấn công chia theo 4 trụ, bấm chạy ngay trên /trust"),
                                    ("0", "lần gọi LLM trên đường quyết định, đọc từ /api/health"),
                                    ("72/72", "ca red-team đạt")]):
        y = 1.3 + i * 1.22
        text(s, 6.3, y, 3.2, 0.55, num, size=30, bold=True, font=HEAD, color=LIME_DEEP if i == 0 else NAVY)
        text(s, 6.3, y + 0.55, 3.2, 0.55, lab, size=12, color=SLATE)

    # 12 ─ Q2: kiến trúc
    s = d.slide(kicker="CÂU HỎI 2 · KIẾN TRÚC", title="Bảy lớp, AI chỉ có mặt ở khâu đọc", notes="""
Câu hỏi hai là kiến trúc. Một chứng từ đi qua bảy lớp: nạp, chuẩn hóa, quyết định tất định, workflow, người trong vòng lặp,
bằng chứng, và vận hành. AI chỉ có mặt ở hai lớp đầu, để đọc ảnh và tra danh tính bên bán, và chỉ để đọc. AI diễn đạt câu hỏi
ở lớp năm thì tắt mặc định. Bên phải là phần thật và phần giả lập: quy tắc, phân quyền, nhật ký, phép thử là mã chạy thật.
Vai người dùng thì đổi được trong phiên demo, danh bạ người gửi là dữ liệu mẫu.""")
    layers = [("Nạp chứng từ", "XML tất định · PDF lớp chữ · ảnh", "AI đọc ảnh"),
              ("Chuẩn hóa, nguồn gốc", "14 trường, kiểm số học, checksum MST", "AI tra MST"),
              ("Quyết định tất định", "3 kiểu dừng, căn cứ theo ngày lập", None),
              ("Workflow", "trạng thái, thẩm quyền, lô, công tắc dừng", None),
              ("Người trong vòng lặp", "câu hỏi “1A, 2A”, hậu kiểm 10%", "AI diễn đạt: tắt"),
              ("Bằng chứng", "chỉ ghi thêm, SHA-256, HMAC, điểm neo", None),
              ("Vận hành", "/api/health, /trust, CI Ubuntu + Windows", None)]
    for i, (name, desc, tag) in enumerate(layers):
        y = 1.3 + i * 0.52
        key = i == 2
        box(s, 0.5, y, 5.75, 0.45, fill=NAVY if key else LIGHT, radius=0.12)
        text(s, 0.62, y, 0.3, 0.45, str(i + 1), size=12, bold=True, color=LIME if key else LIME_DEEP, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 0.92, y, 1.85, 0.45, name, size=12, bold=True, color=WHITE if key else NAVY, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 2.8, y, 2.25 if tag else 3.35, 0.45, desc, size=10.5, color=ON_DARK if key else SLATE, anchor=MSO_ANCHOR.MIDDLE)
        if tag:
            box(s, 5.1, y + 0.08, 1.05, 0.29, fill=AMBER_TINT if "tắt" in tag else LIME_TINT,
                paras=[P(tag, align=PP_ALIGN.CENTER)], size=8.5, bold=True, color=AMBER if "tắt" in tag else LIME_DEEP, pad=0.02)
    box(s, 6.55, 1.3, 2.95, 1.95, fill=LIME_TINT, anchor=MSO_ANCHOR.TOP, pad=0.18,
        paras=[P(R("Thật, không giả lập", bold=True, size=13, color=LIME_DEEP), after=4)] +
              [P(t, bullet=True, after=1) for t in ["Quy tắc, workflow, lô", "Phân quyền kiểm ở server",
                                                     "Nhật ký chống sửa", "12 phép thử trên /trust", "Đọc XML, PDF, ảnh"]],
        size=11.5, color=NAVY)
    box(s, 6.55, 3.4, 2.95, 1.53, fill=AMBER_TINT, anchor=MSO_ANCHOR.TOP, pad=0.18,
        paras=[P(R("Giả lập trong demo", bold=True, size=13, color=AMBER), after=4)] +
              [P(t, bullet=True, after=1, bullet_color=AMBER) for t in
               ["Vai người dùng: đổi vai trong phiên demo", "Danh bạ người gửi là dữ liệu mẫu",
                "LLM dev chỉ bật cho demo"]],
        size=11.5, color=NAVY)

    # 13 ─ Q2: đánh đổi
    s = d.slide(kicker="CÂU HỎI 2 · ĐÁNH ĐỔI", title="Ba đánh đổi đội tự chọn", notes="""
Có ba đánh đổi chúng em tự chọn. Một, quy tắc thay vì để mô hình ngôn ngữ quyết định: được khả năng tái lập và giải trình
với cơ quan thuế, mất độ linh hoạt, từ khóa có thể sót ca lạ. Nếu thêm AI thì AI chỉ được làm chặt hơn, không bao giờ được nới.
Hai, mỗi phiên một file SQLite: cách ly đơn giản, nhưng chỉ chạy trên một máy; cần nhiều instance thì chuyển sang PostgreSQL.
Ba, gọi Vision đồng bộ qua bên thứ ba: nhanh ra demo, nhưng dữ liệu ảnh đi ra nước ngoài; có khách thật thì phải có hợp đồng
xử lý dữ liệu hoặc chạy chế độ chỉ XML.""")
    trade = [("Quy tắc thay vì LLM quyết định", "tái lập, giải trình được với cơ quan thuế.",
              "độ phủ ngữ nghĩa; từ khóa có thể sót ca lạ.", "AI chỉ được làm chặt hơn, không bao giờ được nới."),
             ("Mỗi phiên một file SQLite", "cách ly đơn giản, sao lưu và xóa theo file.",
              "chỉ chạy trên một máy, một writer.", "cần hơn một instance thì chuyển PostgreSQL có RLS."),
             ("Gọi Vision đồng bộ qua bên thứ ba", "đơn giản, nhanh ra demo.",
              "ảnh đi ra nước ngoài; request có thể chờ tới 35 giây.", "có khách thật: hợp đồng xử lý dữ liệu hoặc chỉ XML.")]
    for i, (t, gain, loss, when) in enumerate(trade):
        x = 0.5 + i * 3.1
        box(s, x, 1.3, 2.85, 3.6, fill=LIGHT, radius=0.05)
        text(s, x + 0.22, 1.42, 2.4, 0.5, str(i + 1), size=26, bold=True, font=HEAD, color=LIME_DEEP)
        text(s, x + 0.22, 1.95, 2.45, 0.62, t, size=14, bold=True, color=NAVY)
        text(s, x + 0.22, 2.68, 2.45, 2.15,
             [P(R("Được: ", bold=True, color=LIME_DEEP), gain, after=6),
              P(R("Mất: ", bold=True, color=ROSE), loss, after=6),
              P(R("Khi nào đổi: ", bold=True, color=NAVY), when)], size=13, color=SLATE)

    # 14 ─ Q2: lộ trình
    s = d.slide(kicker="CÂU HỎI 2 · TỪ DEMO LÊN PRODUCTION", title="Lộ trình D0, 30 ngày, 90 ngày", notes="""
Thứ vỡ đầu tiên khi mở rộng không phải bộ quy tắc, vì nó chạy cục bộ. Vỡ trước là khâu đọc ảnh và PDF trên một instance.
Nên lộ trình là: hôm nay đã có ranh giới cho Vision, bánh cóc thẩm quyền, ngân sách AI, trang trust, và lô có người chốt.
Ba mươi ngày tới: đăng nhập một lần và xác thực hai lớp cho CFO, worker đọc file trong container không mạng, hàng đợi,
đưa khóa nhật ký vào KMS. Chín mươi ngày: lấy XML có chữ ký số làm nguồn sự thật và đối chiếu dữ liệu hóa đơn điện tử của cơ quan thuế.
Trên VNG Cloud có thể dùng VKS, vDB PostgreSQL, vStorage có Object Lock, KMS và IAM; phần này cần VNG xác nhận.""")
    phases = [("D0 · đã có", "Vision chỉ 14 trường, bánh cóc thẩm quyền, ngân sách AI, /trust. Hôm nay: lô có người chốt."),
              ("30 ngày", "SSO, MFA cho CFO. Worker đọc file trong container không mạng, không secret. Hàng đợi. Khóa audit vào KMS."),
              ("90 ngày", "XML có chữ ký số là nguồn sự thật, đối chiếu dữ liệu hóa đơn điện tử của cơ quan thuế. PostgreSQL đa tenant.")]
    arrow(s, 0.6, 1.62, 9.4, 1.62, color=LINE, width=2)
    for i, (t, body) in enumerate(phases):
        x = 0.5 + i * 3.1
        circle(s, x + 0.05, 1.47, 0.3, "", fill=LIME_DEEP if i == 0 else NAVY)
        box(s, x + 0.4, 1.42, 1.45, 0.4, fill=WHITE, shape=MSO_SHAPE.RECTANGLE, pad=0.05,
            paras=[P(t)], size=14, bold=True, color=NAVY)
        box(s, x, 1.98, 2.85, 1.75, fill=LIME_TINT if i == 0 else LIGHT, anchor=MSO_ANCHOR.TOP, pad=0.18,
            paras=[P(body)], size=13, color=SLATE)
    box(s, 0.5, 3.92, 9.0, 0.6, fill=WHITE, line=LINE, pad=0.2,
        paras=[P(R("Trên VNG Cloud: ", bold=True, color=NAVY),
                 "VKS + vLB + vWAF · vDB PostgreSQL Cluster · vStorage có Object Lock · KMS · IAM · vMonitor")],
        size=12, color=SLATE)
    text(s, 0.5, 4.62, 9.0, 0.3, "Dịch vụ VNG Cloud lấy từ tài liệu của nhà cung cấp, cần VNG xác nhận trước khi dùng.",
         size=9.5, italic=True, color=MUTED)

    # 15 ─ Riêng tư
    s = d.slide(kicker="RIÊNG TƯ VÀ PHÁP LUẬT", title="Dữ liệu đi đâu, và không đi đâu", notes="""
Về dữ liệu. File XML được đọc bằng mã thường, không bao giờ gửi cho AI. Ảnh và PDF scan thì gửi cho nhà cung cấp Vision,
là bên thứ ba ở nước ngoài, và chỉ lấy về mười bốn trường. Khách hàng nào không muốn thì chạy chế độ chỉ XML.
Mỗi phiên có database riêng, nhật ký chỉ ghi thêm. Về pháp luật, chúng em đã đối chiếu Luật Bảo vệ dữ liệu cá nhân
và Nghị định 356 năm 2025 qua nguồn thứ cấp, nên không trích số điều. Lên production thì cần hợp đồng xử lý dữ liệu với nhà cung cấp AI,
hoặc chạy chỉ XML.""")
    box(s, 0.5, 1.45, 1.35, 0.6, fill=LIGHT, paras=[P("File XML", align=PP_ALIGN.CENTER)], size=13, bold=True, color=NAVY)
    box(s, 0.5, 2.55, 1.35, 0.6, fill=LIGHT, paras=[P("Ảnh, PDF scan", align=PP_ALIGN.CENTER)], size=13, bold=True, color=NAVY, pad=0.04)
    box(s, 2.3, 2.45, 1.75, 0.8, fill=AMBER_TINT, paras=[P(R("Vision AI", bold=True), align=PP_ALIGN.CENTER),
                                                         P("bên thứ ba, ở nước ngoài", align=PP_ALIGN.CENTER, size=10.5)],
        size=13, color=AMBER, pad=0.05)
    box(s, 4.5, 1.45, 1.3, 1.7, fill=NAVY, paras=[P("Bộ quy tắc", align=PP_ALIGN.CENTER)], size=13, bold=True, color=WHITE, pad=0.05)
    arrow(s, 1.9, 1.75, 4.45, 1.75, color=LIME_DEEP, width=1.75)
    text(s, 2.0, 1.4, 2.4, 0.3, "không qua AI", size=11, bold=True, color=LIME_DEEP, align=PP_ALIGN.CENTER)
    arrow(s, 1.9, 2.85, 2.25, 2.85)
    arrow(s, 4.1, 2.85, 4.45, 2.85)
    text(s, 2.3, 3.3, 1.75, 0.3, "chỉ 14 trường", size=11, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    text(s, 0.5, 3.8, 5.3, 1.1, [P(R("Chế độ chỉ XML: ", bold=True, color=NAVY), "không gửi gì cho AI.", after=4),
                                 P("Mỗi phiên một database riêng; nhật ký chỉ ghi thêm.")], size=13, color=SLATE)
    box(s, 6.2, 1.35, 3.3, 3.1, fill=LIGHT, anchor=MSO_ANCHOR.TOP, pad=0.22,
        paras=[P(R("Pháp luật", bold=True, size=14, color=NAVY), after=6),
               P("Luật Bảo vệ dữ liệu cá nhân 91/2025/QH15 và Nghị định 356/2025/NĐ-CP.", after=6),
               P("Đã đối chiếu qua nguồn thứ cấp, nên chưa trích số điều.", after=6),
               P(R("Lên production: ", bold=True, color=NAVY), "hợp đồng xử lý dữ liệu với nhà cung cấp AI, hoặc chạy chỉ XML.")],
        size=13, color=SLATE)

    # 16 ─ Số đo
    s = d.slide(kicker="SỐ ĐO", title="Đo được gì, và đo trên cái gì", notes="""
Đây là các con số đo được, kèm ngữ cảnh. Bộ Verify đạt chín trên chín, gồm bốn ca chung và năm ca Đề A. Năm ca Đề A chạy khoảng
tám mươi mili giây trên bản live. Chạy lại một trăm tám mươi ca hồi quy mất khoảng hai đến ba giây, không quyết định nào bị đổi.
Bảy mươi hai ca red-team đạt cả, sáu mươi ca phát triển đạt cả, mười sáu bộ test tự động. Chúng em còn cho ba vai giám khảo mô phỏng
chấm thử: họ tìm ra mười hai nhóm lỗi làm mất điểm, đội đã sửa và viết test cho từng nhóm.""")
    tiles = [("9/9", "ca Verify đạt: 4 ca chung + 5 ca Đề A"),
             ("~80 ms", "chạy 5 ca Đề A trên bản live"),
             ("0", "quyết định bị đổi khi chạy lại 180 ca hồi quy (khoảng 2–3 giây)"),
             ("72/72", "ca red-team đạt"),
             ("60/60", "ca phát triển đạt"),
             ("16", "bộ test tự động; /trust 12/12 phép thử bảo mật")]
    for i, (num, lab) in enumerate(tiles):
        x = 0.5 + (i % 3) * 3.1
        y = 1.3 + (i // 3) * 1.65
        box(s, x, y, 2.85, 1.5, fill=LIGHT, radius=0.06)
        text(s, x + 0.22, y + 0.12, 2.45, 0.6, num, size=30, bold=True, font=HEAD, color=NAVY)
        text(s, x + 0.22, y + 0.78, 2.45, 0.65, lab, size=12, color=SLATE)
    text(s, 0.5, 4.67, 9.0, 0.3, "Ba vai “giám khảo vòng cuối” mô phỏng tìm ra 12 nhóm lỗi mất điểm; đã sửa, có test.",
         size=11.5, italic=True, color=MUTED)

    # 17 ─ Trước / sau
    s = d.slide(kicker="SỐ ĐO · TRƯỚC VÀ SAU", title="11 ca kiểu giám khảo: bản 13/10 và bản 15/10", notes="""
Còn đây là so sánh trước và sau trên mười một ca viết theo kiểu giám khảo, giữa bản ngày mười ba và bản ngày mười lăm tháng mười.
Số ca xử lý đúng tăng từ ba lên mười một trên mười một. Ca thường quy bị hỏi thừa giảm từ ba trên ba xuống không.
Xếp sai nhóm dừng từ bốn trên tám xuống không. Câu hỏi đạt chất lượng từ một trên mười lên tám trên tám.
Giới hạn em nói luôn: các ca này do đội và AI viết, chưa phải đánh giá độc lập. Buổi thử với ba người dùng ngày mười sáu
chưa cho số liệu định lượng.""")
    compare = [("Xử lý đúng", "3/11", "11/11"), ("Ca thường quy bị hỏi thừa", "3/3", "0/3"),
               ("Xếp sai nhóm dừng", "4/8", "0/8"), ("Câu hỏi đạt chất lượng", "1/10", "8/8")]
    text(s, 5.2, 1.25, 1.5, 0.3, "13/10", size=11, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 7.4, 1.25, 1.8, 0.3, "15/10", size=11, bold=True, color=LIME_DEEP, align=PP_ALIGN.CENTER)
    for i, (lab, before, after) in enumerate(compare):
        y = 1.6 + i * 0.66
        box(s, 0.5, y, 9.0, 0.56, fill=LIGHT if i % 2 == 0 else WHITE, radius=0.1)
        text(s, 0.75, y, 4.3, 0.56, lab, size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 5.2, y, 1.5, 0.56, before, size=20, color=MUTED, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        arrow(s, 6.85, y + 0.28, 7.3, y + 0.28)
        text(s, 7.4, y, 1.8, 0.56, after, size=22, bold=True, font=HEAD, color=LIME_DEEP,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.5, 4.3, 9.0, 0.62, fill=AMBER_TINT, pad=0.2,
        paras=[P(R("Giới hạn: ", bold=True, color=AMBER),
                 "ca do đội và AI viết, chưa phải đánh giá độc lập. Buổi thử 3 người dùng ngày 16/10: chưa có số liệu định lượng.")],
        size=12, color=NAVY)

    # 18 ─ Giới hạn, rủi ro, bước tiếp
    s = d.slide(kicker="NÓI THẲNG", title="Giới hạn, rủi ro và bước tiếp theo", notes="""
Phần này chúng em muốn nói thẳng. Giới hạn: vai người dùng trong demo là mô phỏng, trình đọc PDF chưa chạy trong container riêng,
khóa ký nhật ký nằm cùng máy, và SQLite chỉ chạy trên một máy. Rủi ro lớn nhất: ảnh bị đọc sai số tiền mà các con số vẫn khớp nhau,
thì ca có thể đi luồng thấp hơn. Dữ kiện như ủy nhiệm chi vẫn do người khai. Bước tiếp theo: đăng nhập một lần và xác thực hai lớp,
worker đọc file cô lập, đối chiếu dữ liệu hóa đơn điện tử của cơ quan thuế, và thử với người dùng thật để có số liệu.""")
    cols3 = [("Giới hạn", AMBER, AMBER_TINT, ["Vai người dùng trong demo là mô phỏng", "Trình đọc PDF chưa ở container riêng",
                                               "Khóa HMAC nằm cùng máy", "SQLite chỉ chạy trên một máy"]),
             ("Rủi ro", ROSE, ROSE_TINT, ["Ảnh đọc sai số tiền mà vẫn cân đối: ca đi luồng thấp hơn",
                                          "Dữ kiện do người khai, như ủy nhiệm chi", "Ảnh gửi cho Vision ở nước ngoài"]),
             ("Bước tiếp", LIME_DEEP, LIME_TINT, ["SSO, MFA cho CFO", "Worker đọc file cô lập",
                                                  "Đối chiếu dữ liệu hóa đơn điện tử của cơ quan thuế",
                                                  "Thử với người dùng thật, có số liệu"])]
    for i, (t, c, fill, items) in enumerate(cols3):
        x = 0.5 + i * 3.1
        box(s, x, 1.3, 2.85, 3.6, fill=fill, anchor=MSO_ANCHOR.TOP, pad=0.22,
            paras=[P(R(t, bold=True, size=16, color=c), after=8)] +
                  [P(it, bullet=True, bullet_color=c, after=8) for it in items],
            size=14, color=NAVY)

    # 19 ─ Demo
    s = d.slide(dark=True, kicker="DEMO TRỰC TIẾP", title="Sáu bước trên bản đang chạy", notes="""
Giờ em demo trên bản đang chạy. Kế toán viên tải một lô nhiều file. Ban giám khảo sẽ thấy lô được xếp theo mã số thuế, ngày, số,
và file trùng chỉ tính một lần. Em mở phiếu giải trình của một ca được duyệt và một ca bị dừng tô đỏ. Kế toán trưởng trả lời
nhóm câu hỏi một lần. Người tạo lô thử tự chốt lô của mình và bị chặn, rồi người có quyền chốt lô. Cuối cùng là màn hình tổng quan
cho sếp và trang trust chạy mười hai phép thử. Nếu mạng hội trường chậm, em chiếu video một phút.""")
    demo = ["Kế toán viên tải một lô nhiều file",
            "Lô xếp MST → ngày → số; file trùng chỉ tính một",
            "Mở phiếu giải trình: một ca duyệt, một ca dừng tô đỏ",
            "Kế toán trưởng trả lời nhóm câu hỏi “1A, 2A”",
            "Người tạo lô tự chốt: bị chặn. Người có quyền chốt lô",
            "Xem tổng quan cho sếp, chạy /trust 12/12"]
    for i, st in enumerate(demo):
        y = 1.32 + i * 0.6
        circle(s, 0.6, y + 0.04, 0.4, str(i + 1), fill=LIME, color=NAVY, size=13)
        text(s, 1.2, y, 5.0, 0.48, st, size=14, color=ON_DARK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 6.6, 1.32, 2.9, 2.0, fill=CARD_DARK, anchor=MSO_ANCHOR.MIDDLE, pad=0.25,
        paras=[P(R("Video 1 phút", bold=True, size=20, color=LIME, font=HEAD), after=6),
               P("Dự phòng khi mạng hội trường chậm.")], size=13, color=LINE)
    text(s, 6.6, 3.55, 2.9, 0.9, [P("Bản đang chạy", size=11, bold=True, color=LIME, after=2),
                                  P("limegreen-salamander-517622", size=11, color=LINE),
                                  P(".hostingersite.com", size=11, color=LINE)])

    out.parent.mkdir(parents=True, exist_ok=True)
    d.prs.save(str(out))
    return d


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--screens", type=Path, default=ROOT / "docs" / "submission" / "screens")
    ap.add_argument("--out", type=Path, default=ROOT / "docs" / "submission" / "Tax-Referee-trinh-bay-cai-tien-17-10.pptx")
    a = ap.parse_args()
    d = build(a.screens, a.out)
    print(f"Đã ghi {a.out} · {d.n} slide")
    print("Ảnh đã chèn: " + (", ".join(d.inserted) or "không có"))
    print("Khung chờ: " + (", ".join(d.placeholders) or "không có"))


if __name__ == "__main__":
    os.environ.setdefault("PYTHONIOENCODING", "utf-8")
    main()
